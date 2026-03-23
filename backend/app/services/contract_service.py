"""
合同生成服务
"""
import os
import uuid
from datetime import datetime
from typing import Optional, List, Dict, Any
from sqlalchemy.orm import Session

from app.models.document import Template, FieldDefinition, Contract
from app.services.document_service import DocumentParser, TemplateRenderer
from app.config import settings


class ContractService:
    """合同服务"""

    @staticmethod
    def generate_contract_no() -> str:
        """生成合同编号"""
        # 格式: HT-YYYYMMDD-XXXX
        date_str = datetime.now().strftime("%Y%m%d")
        random_str = str(uuid.uuid4())[:4].upper()
        return f"HT-{date_str}-{random_str}"

    @staticmethod
    async def get_field_definition(
        contract_type: str,
        db: Session
    ) -> Optional[FieldDefinition]:
        """获取合同类型的字段定义"""
        return db.query(FieldDefinition).filter(
            FieldDefinition.contract_type == contract_type,
            FieldDefinition.is_active == True
        ).first()

    @staticmethod
    async def get_template(
        contract_type: str,
        db: Session
    ) -> Optional[Template]:
        """获取合同类型的模板"""
        return db.query(Template).filter(
            Template.contract_type == contract_type,
            Template.is_active == True
        ).first()

    @staticmethod
    def parse_field_definition(fields_json: str) -> List[Dict[str, Any]]:
        """解析字段定义 JSON"""
        import json
        return json.loads(fields_json)

    @staticmethod
    def get_required_fields(field_definition: FieldDefinition) -> List[Dict[str, Any]]:
        """获取必填字段列表"""
        fields = ContractService.parse_field_definition(field_definition.fields)
        return [f for f in fields if f.get("required", True)]

    @staticmethod
    def validate_fields(
        field_definition: FieldDefinition,
        field_values: Dict[str, Any]
    ) -> List[str]:
        """
        验证字段值

        Returns:
            缺失的必填字段列表
        """
        required_fields = ContractService.get_required_fields(field_definition)
        missing = []

        for field in required_fields:
            name = field.get("name")
            if name not in field_values or not field_values[name]:
                missing.append(field.get("label", name))

        return missing

    @staticmethod
    def check_field_completion(
        field_definition: FieldDefinition,
        field_values: Dict[str, Any]
    ) -> Dict[str, Any]:
        """
        检查字段完成状态

        Returns:
            {
                "total": 总字段数,
                "filled": 已填字段数,
                "missing_fields": 缺失字段列表,
                "completion_rate": 完成率
            }
        """
        fields = ContractService.parse_field_definition(field_definition.fields)
        total = len(fields)
        filled = 0
        missing_fields = []

        for field in fields:
            name = field.get("name")
            if name in field_values and field_values[name]:
                filled += 1
            elif field.get("required", True):
                missing_fields.append(field.get("label", name))

        completion_rate = filled / total if total > 0 else 0

        return {
            "total": total,
            "filled": filled,
            "missing_fields": missing_fields,
            "completion_rate": round(completion_rate, 2)
        }

    @staticmethod
    async def generate_contract_content(
        template: Template,
        field_values: Dict[str, Any]
    ) -> str:
        """
        生成合同内容

        Args:
            template: 合同模板
            field_values: 字段值

        Returns:
            生成的合同内容
        """
        # 如果模板有内容，直接使用
        if template.content:
            return TemplateRenderer.render_simple(template.content, field_values)

        # 否则从文件解析
        content, _ = await DocumentParser.parse_file(template.file_path)
        return TemplateRenderer.render_simple(content, field_values)

    @staticmethod
    async def create_contract(
        user_id: int,
        contract_type: str,
        field_values: Dict[str, Any],
        db: Session,
        template_id: Optional[int] = None,
        session_id: Optional[int] = None
    ) -> Contract:
        """
        创建合同记录

        Args:
            user_id: 用户ID
            contract_type: 合同类型
            field_values: 字段值
            db: 数据库会话
            template_id: 模板ID
            session_id: 会话ID

        Returns:
            合同记录
        """
        import json

        # 获取模板
        template = None
        if template_id:
            template = db.query(Template).filter(Template.id == template_id).first()
        else:
            template = await ContractService.get_template(contract_type, db)

        # 生成合同内容
        content = ""
        if template:
            content = await ContractService.generate_contract_content(template, field_values)

        # 创建合同记录
        contract = Contract(
            contract_no=ContractService.generate_contract_no(),
            user_id=user_id,
            template_id=template.id if template else None,
            contract_type=contract_type,
            title=f"{contract_type}-{datetime.now().strftime('%Y%m%d')}",
            status="draft",
            field_values=json.dumps(field_values, ensure_ascii=False),
            content=content,
            session_id=session_id
        )

        db.add(contract)
        db.commit()
        db.refresh(contract)

        return contract

    @staticmethod
    async def update_contract(
        contract_id: int,
        field_values: Dict[str, Any],
        db: Session
    ) -> Contract:
        """更新合同字段值"""
        import json

        contract = db.query(Contract).filter(Contract.id == contract_id).first()
        if not contract:
            raise ValueError("合同不存在")

        # 更新字段值
        contract.field_values = json.dumps(field_values, ensure_ascii=False)

        # 重新生成内容
        if contract.template_id:
            template = db.query(Template).filter(Template.id == contract.template_id).first()
            if template:
                contract.content = await ContractService.generate_contract_content(template, field_values)

        db.commit()
        db.refresh(contract)

        return contract

    @staticmethod
    async def export_to_pdf(
        contract: Contract,
        output_path: Optional[str] = None
    ) -> str:
        """
        导出合同为 PDF

        Args:
            contract: 合同记录
            output_path: 输出路径（可选）

        Returns:
            PDF 文件路径
        """
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import cm
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

        # 生成输出路径
        if not output_path:
            output_path = os.path.join(
                settings.CONTRACTS_DIR,
                f"{contract.contract_no}.pdf"
            )

        # 注册中文字体（需要系统有中文字体）
        try:
            # Windows 系统字体路径
            font_path = "C:/Windows/Fonts/simsun.ttc"
            if os.path.exists(font_path):
                pdfmetrics.registerFont(TTFont('SimSun', font_path))
                chinese_font = 'SimSun'
            else:
                chinese_font = 'Helvetica'
        except:
            chinese_font = 'Helvetica'

        # 创建 PDF
        doc = SimpleDocTemplate(output_path, pagesize=A4)
        styles = getSampleStyleSheet()

        # 创建中文样式
        chinese_style = ParagraphStyle(
            'Chinese',
            parent=styles['Normal'],
            fontName=chinese_font,
            fontSize=12,
            leading=20
        )

        title_style = ParagraphStyle(
            'ChineseTitle',
            parent=styles['Title'],
            fontName=chinese_font,
            fontSize=18,
            leading=30,
            alignment=1  # 居中
        )

        # 构建文档内容
        story = []

        # 标题
        story.append(Paragraph(contract.title or contract.contract_no, title_style))
        story.append(Spacer(1, 1 * cm))

        # 合同编号
        story.append(Paragraph(f"合同编号：{contract.contract_no}", chinese_style))
        story.append(Spacer(1, 0.5 * cm))

        # 合同内容
        if contract.content:
            for line in contract.content.split("\n"):
                if line.strip():
                    story.append(Paragraph(line, chinese_style))
                    story.append(Spacer(1, 0.3 * cm))

        # 签名区域
        story.append(Spacer(1, 2 * cm))
        story.append(Paragraph("_" * 20 + "    " + "_" * 20, chinese_style))
        story.append(Paragraph("甲方签字              乙方签字", chinese_style))
        story.append(Spacer(1, 1 * cm))
        story.append(Paragraph(f"日期：{datetime.now().strftime('%Y年%m月%d日')}", chinese_style))

        # 生成 PDF
        doc.build(story)

        return output_path

    @staticmethod
    async def export_to_docx(
        contract: Contract,
        output_path: Optional[str] = None
    ) -> str:
        """
        导出合同为 DOCX

        Args:
            contract: 合同记录
            output_path: 输出路径（可选）

        Returns:
            DOCX 文件路径
        """
        from docx import Document
        from docx.shared import Pt, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        # 生成输出路径
        if not output_path:
            output_path = os.path.join(
                settings.CONTRACTS_DIR,
                f"{contract.contract_no}.docx"
            )

        # 创建文档
        doc = Document()

        # 设置标题
        title = doc.add_heading(contract.title or contract.contract_no, 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 合同编号
        p = doc.add_paragraph()
        p.add_run(f"合同编号：{contract.contract_no}").bold = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()

        # 合同内容
        if contract.content:
            for line in contract.content.split("\n"):
                if line.strip():
                    p = doc.add_paragraph(line)

        # 签名区域
        doc.add_paragraph()
        doc.add_paragraph()

        # 添加签名表格
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "甲方签字："
        table.cell(0, 1).text = "乙方签字："
        table.cell(1, 0).text = "_" * 15
        table.cell(1, 1).text = "_" * 15

        doc.add_paragraph()
        doc.add_paragraph(f"日期：{datetime.now().strftime('%Y年%m月%d日')}")

        # 保存文档
        doc.save(output_path)

        return output_path
