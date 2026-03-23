"""
文档解析服务 - 解析各种格式的文档
"""
import os
import re
from typing import Optional, Tuple, List
from pathlib import Path

from app.config import settings


class DocumentParser:
    """文档解析器"""

    @staticmethod
    async def parse_file(file_path: str) -> Tuple[str, dict]:
        """
        解析文件，返回文本内容和元数据

        Args:
            file_path: 文件路径

        Returns:
            (文本内容, 元数据)
        """
        ext = Path(file_path).suffix.lower()

        if ext == ".txt":
            return await DocumentParser.parse_txt(file_path)
        elif ext == ".md":
            return await DocumentParser.parse_markdown(file_path)
        elif ext == ".pdf":
            return await DocumentParser.parse_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            return await DocumentParser.parse_docx(file_path)
        else:
            raise ValueError(f"不支持的文件类型: {ext}")

    @staticmethod
    async def parse_txt(file_path: str) -> Tuple[str, dict]:
        """解析 TXT 文件"""
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()

        # 统计信息
        lines = content.split("\n")
        words = len(content)

        metadata = {
            "type": "txt",
            "lines": len(lines),
            "characters": words
        }

        return content, metadata

    @staticmethod
    async def parse_markdown(file_path: str) -> Tuple[str, dict]:
        """解析 Markdown 文件"""
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()

        # 提取标题
        titles = re.findall(r"^#+\s+(.+)$", content, re.MULTILINE)

        # 提取代码块数量
        code_blocks = len(re.findall(r"```", content)) // 2

        metadata = {
            "type": "markdown",
            "titles": titles[:10],  # 前10个标题
            "code_blocks": code_blocks,
            "characters": len(content)
        }

        return content, metadata

    @staticmethod
    async def parse_pdf(file_path: str) -> Tuple[str, dict]:
        """解析 PDF 文件"""
        try:
            import pypdf

            with open(file_path, "rb") as f:
                reader = pypdf.PdfReader(f)
                pages = len(reader.pages)

                text_parts = []
                for page in reader.pages:
                    text = page.extract_text()
                    if text:
                        text_parts.append(text)

                content = "\n\n".join(text_parts)

                metadata = {
                    "type": "pdf",
                    "pages": pages,
                    "characters": len(content)
                }

                return content, metadata

        except ImportError:
            raise ImportError("请安装 pypdf: pip install pypdf")

    @staticmethod
    async def parse_docx(file_path: str) -> Tuple[str, dict]:
        """解析 DOCX 文件"""
        try:
            from docx import Document

            doc = Document(file_path)

            # 提取段落文本
            paragraphs = []
            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text.strip())

            # 提取表格文本
            tables_text = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    if row_text.strip():
                        tables_text.append(row_text)

            content = "\n\n".join(paragraphs)
            if tables_text:
                content += "\n\n" + "\n".join(tables_text)

            metadata = {
                "type": "docx",
                "paragraphs": len(paragraphs),
                "tables": len(doc.tables),
                "characters": len(content)
            }

            return content, metadata

        except ImportError:
            raise ImportError("请安装 python-docx: pip install python-docx")

    @staticmethod
    def extract_variables_from_template(content: str) -> List[str]:
        """
        从模板内容中提取变量

        支持的变量格式：
        - {{variable_name}}
        - {variable_name}
        - 【变量名】

        Returns:
            变量名列表
        """
        variables = set()

        # Jinja2 格式 {{ variable }}
        jinja_vars = re.findall(r"\{\{\s*(\w+)\s*\}\}", content)
        variables.update(jinja_vars)

        # Python 格式 { variable }
        python_vars = re.findall(r"\{(\w+)\}", content)
        variables.update(python_vars)

        # 中文格式 【变量名】
        chinese_vars = re.findall(r"【(.+?)】", content)
        variables.update(chinese_vars)

        return list(variables)

    @staticmethod
    def clean_text(text: str) -> str:
        """清理文本（去除多余空白、特殊字符等）"""
        # 去除多余空白
        text = re.sub(r"\s+", " ", text)

        # 去除特殊控制字符
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]", "", text)

        return text.strip()

    @staticmethod
    def split_by_sections(content: str) -> List[Tuple[str, str]]:
        """
        按章节分割文档

        Returns:
            [(章节标题, 章节内容)]
        """
        # 匹配常见的章节标题格式
        section_pattern = r"^(第[一二三四五六七八九十\d]+[章节条款]|[一二三四五六七八九十]+[、.]|\d+[、.]|\(\d+\))\s*(.+)$"

        sections = []
        current_title = "开头"
        current_content = []

        for line in content.split("\n"):
            match = re.match(section_pattern, line.strip())
            if match:
                # 保存上一个章节
                if current_content:
                    sections.append((current_title, "\n".join(current_content)))
                current_title = match.group(0).strip()
                current_content = []
            else:
                current_content.append(line)

        # 保存最后一个章节
        if current_content:
            sections.append((current_title, "\n".join(current_content)))

        return sections


class TemplateRenderer:
    """模板渲染器"""

    @staticmethod
    def render(template_content: str, variables: dict) -> str:
        """
        渲染模板

        Args:
            template_content: 模板内容
            variables: 变量字典

        Returns:
            渲染后的内容
        """
        from jinja2 import Template

        template = Template(template_content)
        return template.render(**variables)

    @staticmethod
    def render_simple(template_content: str, variables: dict) -> str:
        """
        简单渲染（使用字符串替换）

        支持：
        - {{variable}}
        - {variable}
        - 【变量名】
        """
        result = template_content

        for key, value in variables.items():
            # Jinja2 格式
            result = result.replace(f"{{{{{key}}}}}", str(value))
            # Python 格式
            result = result.replace(f"{{{key}}}", str(value))
            # 中文格式
            result = result.replace(f"【{key}】", str(value))

        return result
